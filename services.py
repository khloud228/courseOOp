import openpyxl

from docx import Document


class Parser():
    def __init__(self, filename:str):
        """
        Инициализация
        """
        self.filename = filename


    def set_sheet(self):
        """
        Создаёт новую страницу
        """
    

    def save(self):
        """
        Выполняет сохранение экзель файла
        """


    def transfer(self):
        """
        Собирает данные из Word таблиц и погружает их в отдельную страницу файла
        """


class WordParser(Parser):
    def __init__(self, filename:str='untitle') -> None:
        super().__init__(filename=filename)
        self.wordDoc = Document(f'{self.filename}.docx')
        self.wBook = openpyxl.Workbook(f'{self.filename}.xlsx')
        self.label = 1


    def set_sheet(self) -> None:
        self.sheet = self.wBook.create_sheet(f"Sheet number of {self.label}")


    def save(self) -> None:
        self.wBook.save(f'{self.filename}.xlsx')


    def transfer(self) -> None:
        for table in self.wordDoc.tables:
            self.set_sheet()
            for row in table.rows:
                rowitems = []
                for cell in row.cells:
                    rowitems.append(cell.text)
                self.sheet.append(rowitems)
            self.label += 1
        self.save()